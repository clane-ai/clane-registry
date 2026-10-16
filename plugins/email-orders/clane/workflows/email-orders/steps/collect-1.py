def main():
    import email, glob, html, json, os, re
    from email import policy

    MAX_TEXT = 60000
    MIN_PAGE_CHARS = 40
    MIN_IMAGE_BYTES = 200000
    SOURCES = "order-sources.md"
    DOC_EXT = (".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".csv", ".txt", ".xlsx", ".docx", ".htm", ".html")
    IGNORE = {"payload.json"}

    def html_to_text(raw):
        raw = re.sub(r"(?is)<(script|style).*?</\1>", "", raw)
        raw = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</tr>|</li>", "\n", raw)
        raw = re.sub(r"(?i)</t[dh]>", " | ", raw)
        return html.unescape(re.sub(r"<[^>]+>", "", raw))

    def tidy(text):
        text = re.sub(r"[ \t\u00a0\u200b]+", " ", (text or "").replace("\r", ""))
        return re.sub(r"\n\s*\n+", "\n", text).strip()

    def safe_name(name, taken):
        base = re.sub(r"[^A-Za-z0-9._-]+", "_", os.path.basename(name or "attachment")).strip("._") or "attachment"
        stem, ext = os.path.splitext(base)
        out, n = base, 1
        while out in taken or os.path.exists(out):
            n += 1
            out = "%s-%d%s" % (stem, n, ext)
        taken.add(out)
        return out

    def email_text_from_trigger(raw):
        raw = raw or ""
        try:
            data = json.loads(raw)
        except ValueError:
            return raw
        if isinstance(data, dict):
            parts = [str(data[k]) for k in ("from", "subject", "email", "body", "text") if data.get(k)]
            return "\n".join(parts) if parts else ""
        return ""

    emails, extracted, taken = [], [], set()
    trigger_text = tidy(email_text_from_trigger(getattr(trigger, "output", "")))
    if trigger_text:
        emails.append({"source": "email body", "text": trigger_text[:MAX_TEXT]})

    for path in sorted(glob.glob("*.eml")):
        with open(path, "rb") as fh:
            msg = email.message_from_bytes(fh.read(), policy=policy.default)
        body = msg.get_body(preferencelist=("plain", "html"))
        text = ""
        if body is not None:
            text = body.get_content()
            if body.get_content_type() == "text/html":
                text = html_to_text(text)
        head = "From: %s\nTo: %s\nSubject: %s\nDate: %s\n" % (msg["From"], msg["To"], msg["Subject"], msg["Date"])
        emails.append({"source": "email " + path, "text": (head + tidy(text))[:MAX_TEXT]})
        for part in msg.walk():
            name = part.get_filename()
            if not name or not name.lower().endswith(DOC_EXT):
                continue
            payload = part.get_payload(decode=True) or b""
            is_image = name.lower().endswith((".png", ".jpg", ".jpeg", ".tif", ".tiff"))
            if not payload or (is_image and len(payload) < MIN_IMAGE_BYTES):
                continue
            out = safe_name(name, taken)
            with open(out, "wb") as fh:
                fh.write(payload)
            extracted.append(out)

    documents, skipped = [], []
    for path in sorted(set(glob.glob("*")) - set(glob.glob("*.eml"))):
        low = path.lower()
        if not os.path.isfile(path) or path in IGNORE or low.endswith(".ocr.md"):
            continue
        if not low.endswith(DOC_EXT):
            skipped.append({"file": path, "why": "not a document type this reads"})
            continue
        size = os.path.getsize(path)
        doc = {"file": path, "text": "", "needs_ocr": False}
        if low.endswith((".png", ".jpg", ".jpeg", ".tif", ".tiff")):
            if size < MIN_IMAGE_BYTES:
                skipped.append({"file": path, "why": "small image (logo or signature)"})
                continue
            doc["needs_ocr"] = True
        elif low.endswith(".pdf"):
            import pdfplumber
            with pdfplumber.open(path) as pdf:
                pages = [(pg.extract_text() or "") for pg in pdf.pages]
            text = "\n".join(pages)
            doc["pages"] = len(pages)
            if len(text.strip()) < MIN_PAGE_CHARS * max(1, len(pages)):
                doc["needs_ocr"] = True
            else:
                doc["text"] = tidy(text)[:MAX_TEXT]
        elif low.endswith((".htm", ".html")):
            with open(path, encoding="utf-8", errors="replace") as fh:
                doc["text"] = tidy(html_to_text(fh.read()))[:MAX_TEXT]
        elif low.endswith((".csv", ".txt")):
            with open(path, encoding="utf-8-sig", errors="replace") as fh:
                doc["text"] = tidy(fh.read())[:MAX_TEXT]
        elif low.endswith(".xlsx"):
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            rows = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    cells = ["" if c is None else str(c) for c in row]
                    if any(cells):
                        rows.append(" | ".join(cells))
            doc["text"] = "\n".join(rows)[:MAX_TEXT]
        elif low.endswith(".docx"):
            import docx
            doc["text"] = tidy("\n".join(p.text for p in docx.Document(path).paragraphs))[:MAX_TEXT]
        documents.append(doc)

    ocr = [d["file"] for d in documents if d["needs_ocr"]]
    if not emails and not documents:
        raise ValueError("Nothing to read: the email had no text and no documents.")
    parts = ["## Email (%s)\n\n%s" % (e["source"], e["text"]) for e in emails]
    for d in documents:
        body = "(scanned: read it with ocr_document)" if d["needs_ocr"] else d["text"]
        parts.append("## Document %s\n\n%s" % (d["file"], body))
    with open(SOURCES, "w", encoding="utf-8", newline="\n") as fh:
        fh.write("\n\n".join(parts) + "\n")
    return {
        "output": "Read %d email%s and %d document%s; %d need OCR. Everything is in %s." % (
            len(emails), "" if len(emails) == 1 else "s",
            len(documents), "" if len(documents) == 1 else "s", len(ocr), SOURCES),
        "sources_file": SOURCES,
        "ocr_files": ocr,
        "documents": [{"file": d["file"], "needs_ocr": d["needs_ocr"], "chars": len(d["text"])}
                      for d in documents],
        "skipped": skipped,
    }
