def main():
    import csv, glob, json, re
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, Side

    raw = extracted.output or ""
    # The role may wrap its answer in thinking tags or ```json fences - strip
    # both, then grab the JSON object. Fail with a short, readable message (not
    # a giant traceback) when the model returned no JSON at all.
    raw = re.sub(r"<thinking>.*?</thinking>", "", raw, flags=re.S | re.I)
    raw = raw.replace("```json", "").replace("```", "")
    m = re.search(r"\{.*\}", raw, re.S)
    if not m:
        raise ValueError("extraction returned no JSON object; got: " + raw.strip()[:200])
    spec = json.loads(m.group(0))

    def norm(s):
        return re.sub(r"\s+", "", (s or "")).upper()

    def find_csv(columns):
        # Bound setup files are staged into the working folder; find each one
        # by its header rather than by its file name.
        for path in sorted(glob.glob("*.csv")):
            with open(path, encoding="utf-8-sig", newline="") as fh:
                header = {h.strip() for h in next(csv.reader(fh), [])}
            if set(columns) <= header:
                return path
        return None

    # 1) Resolve stock codes and the customer account code from the two setup
    #    files (code_mapping, customer_registry) when they are bound.
    notes = []
    mapping = find_csv(["customer", "customer_code", "stock_code"])
    if mapping:
        rows = list(csv.DictReader(open(mapping, encoding="utf-8-sig")))
        by_cust, by_stock = {}, {}
        for r in rows:
            cc, sc = norm(r["customer_code"]), norm(r["stock_code"])
            if cc:
                by_cust.setdefault(cc, r)
            if sc:
                by_stock.setdefault(sc, r)
        keys = sorted({r["customer"] for r in rows if r["customer"]})

        mapped, unmapped, votes = 0, [], {}
        for i, ln in enumerate(spec.get("lines", []), 1):
            code = norm(ln.get("code"))
            if not code:
                unmapped.append("L%d:(blank)" % i)
                continue
            cust = by_cust.get(code)
            row = cust or by_stock.get(code)
            if row:
                ln["code"] = (row["stock_code"] or "").strip()
                mapped += 1
                # Only a DISTINCT customer code identifies the customer
                # (pass-through codes where customer code == stock code do not).
                if cust and norm(cust["customer_code"]) != norm(cust["stock_code"]):
                    votes[row["customer"]] = votes.get(row["customer"], 0) + 1
            else:
                unmapped.append("L%d:%s" % (i, ln.get("code")))

        cust_in = spec.get("customer")
        key = (max(votes, key=votes.get) if votes
               else next((c for c in keys if norm(c) in norm(cust_in) or norm(cust_in) in norm(c)),
                         cust_in))

        acct = None
        registry = find_csv(["customer_name", "account_code"])
        if registry:
            reg = {norm(r.get("customer_name")): (r.get("account_code") or "").strip()
                   for r in csv.DictReader(open(registry, encoding="utf-8-sig"))}
            acct = reg.get(norm(key)) or None
        spec["customer"] = acct or key

        notes.append("Mapped %d/%d lines via the code mapping." % (mapped, len(spec.get("lines", []))))
        if unmapped:
            notes.append("Unmapped (add a mapping row): " + ", ".join(unmapped) + ".")
        if not acct:
            notes.append("Account code for '%s' not set in the customer registry." % key)

    if notes:
        spec["note"] = " ".join(notes) + ((" | " + spec["note"]) if spec.get("note") else "")

    # 2) Render the Order Form.
    PRICE = '#,##0.00_ ;\\-#,##0.00\\ '
    BLUE, PO_BLUE = "FF0070C0", "FF1F4E79"
    s = Side(style="thin")
    BORDER = Border(left=s, right=s, top=s, bottom=s)

    def F(bold=False, color=None, name="Calibri", size=10):
        return Font(name=name, size=size, bold=bold, color=color)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Order Form"
    for col, w in {"A": 17.0, "B": 52.0, "C": 12.5703125, "D": 13.7109375}.items():
        ws.column_dimensions[col].width = w

    def put(coord, val, font):
        ws[coord] = val
        ws[coord].font = font

    put("B3", "Customer:", F());              put("C3", spec.get("customer"), F())
    put("B4", "Store/Branch:", F());          put("C4", spec.get("store_branch"), F())
    put("B5", "Required Date:", F());         put("C5", spec.get("required_date"), F())
    put("B6", "Customer Order Number:", F())
    put("C6", spec.get("po_number"), F(name="Century Gothic", size=11, color=PO_BLUE))
    if spec.get("note"):
        put("J4", spec["note"], F(name="Arial"))

    put("A11", "Order Qty.", F(bold=True))
    put("B11", "Stock Code", F(bold=True))
    put("C11", "Price", F(bold=True))
    for i, ln in enumerate(spec.get("lines", [])):
        r = 12 + i
        a = ws.cell(row=r, column=1, value=ln.get("qty"))
        a.font = F(bold=True); a.alignment = Alignment(horizontal="left"); a.border = BORDER
        b = ws.cell(row=r, column=2, value=ln.get("code"))
        b.font = F(bold=True, color=BLUE); b.number_format = "@"; b.border = BORDER
        c = ws.cell(row=r, column=3, value=ln.get("price"))
        c.font = F(bold=True, color=BLUE); c.number_format = PRICE
        c.alignment = Alignment(horizontal="left"); c.border = BORDER

    po = str(spec.get("po_number") or "order_form").strip().replace("/", "_")
    out = po + ".xlsx"
    wb.save(out)
    n = len(spec.get("lines", []))
    cust = spec.get("customer") or "the customer"
    return {"file": out, "output": "Order Form ready for %s - PO %s, %d line%s. Saved as %s." % (
        cust, spec.get("po_number") or "(none)", n, "" if n == 1 else "s", out)}
